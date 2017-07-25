import os
from xml.dom import minidom

from openpyxl import load_workbook


def read_file(path):
   wb=load_workbook(path,True)
   ws=wb['student']
   list=[]
   r=0;c=0;
   for row in ws.rows:
       for cell in row:
           list.append(cell.value)
           c=++c
       r=++r
   return list

def create_xml(list,output_path):
    doc=minidom.Document()
    root=doc.createElement('students')

    for i in  range(int(len(list)/5)):
        student=doc.createElement('student')

        node_index=doc.createElement('index')
        node_index_value=doc.createTextNode(str(list[i*5]))
        node_index.appendChild(node_index_value)

        node_name=doc.createElement('name')
        node_name_value=doc.createTextNode(str(list[i*5+1]))
        node_name.appendChild(node_name_value)

        node_mathematics=doc.createElement('mathematics')
        node_mathematics_value=doc.createTextNode(str(list[i*5+2]))
        node_mathematics.appendChild(node_mathematics_value)

        node_physical=doc.createElement('physical')
        node_physical_value=doc.createTextNode(str(list[i*5+3]))
        node_physical.appendChild(node_physical_value)

        node_Chemistry=doc.createElement('Chemistry')
        node_Chemistry_value=doc.createTextNode(str(list[i*5+4]))
        node_Chemistry.appendChild(node_Chemistry_value)

        student.appendChild(node_index)
        student.appendChild(node_name)
        student.appendChild(node_mathematics)
        student.appendChild(node_physical)
        student.appendChild(node_Chemistry)
        root.appendChild(student)
    doc.appendChild(root)

    with open(output_path,'w',encoding='utf-8') as f:
        doc.writexml(f,indent='\t', addindent='\t', newl='\n',encoding='utf-8')


if __name__=='__main__':
    input_path=os.getcwd()+os.path.sep+"students.xlsx"
    output_path=os.getcwd()+os.path.sep+'students.xml'
    list=read_file(input_path)
    create_xml(list,output_path)
